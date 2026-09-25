"""
reports.py
----------
Generates downloadable reports: an individual student's PDF report card,
and a class-wide Excel summary. Both pull live data from the database.
"""

import os
import sys
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.enums import TA_CENTER

from database import get_marks_for_student, get_attendance_percentage, get_overall_percentage
from ml_model import predict_risk
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

if getattr(sys, 'frozen', False):
    APP_DATA_DIR = os.path.join(os.path.expanduser("~"), "StudentAnalyticsSystem")
    os.makedirs(APP_DATA_DIR, exist_ok=True)
    REPORTS_DIR = os.path.join(APP_DATA_DIR, "reports_output")
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    REPORTS_DIR = os.path.join(BASE_DIR, "reports_output")


def generate_student_report_pdf(student):
    """student is a row from get_all_students() / search_students() -
    a dict-like object with id, name, roll_no, class_name, section."""
    os.makedirs(REPORTS_DIR, exist_ok=True)
    filename = os.path.join(REPORTS_DIR, f"report_{student['roll_no']}.pdf")

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title", parent=styles["Title"], fontSize=18, alignment=TA_CENTER
    )

    story = []
    story.append(Paragraph("Student Report Card", title_style))
    story.append(Spacer(1, 0.5 * cm))

    info_text = (
        f"<b>Name:</b> {student['name']}<br/>"
        f"<b>Roll No:</b> {student['roll_no']}<br/>"
        f"<b>Class:</b> {student['class_name']} - {student['section']}"
    )
    story.append(Paragraph(info_text, styles["Normal"]))
    story.append(Spacer(1, 0.5 * cm))

    attendance_pct = get_attendance_percentage(student["id"])
    marks_pct = get_overall_percentage(student["id"])
    risk = predict_risk(attendance_pct, marks_pct)

    summary_text = (
        f"<b>Overall Attendance:</b> {attendance_pct}%<br/>"
        f"<b>Overall Marks:</b> {marks_pct}%<br/>"
        f"<b>Status:</b> {risk}"
    )
    story.append(Paragraph(summary_text, styles["Normal"]))
    story.append(Spacer(1, 0.5 * cm))

    # --- Marks breakdown table ---
    story.append(Paragraph("Marks Breakdown", styles["Heading2"]))
    marks_rows = get_marks_for_student(student["id"])

    table_data = [["Subject", "Exam", "Marks Obtained", "Max Marks"]]
    for row in marks_rows:
        table_data.append([
            row["subject_name"], row["exam_type"],
            str(row["marks_obtained"]), str(row["max_marks"])
        ])

    if len(table_data) == 1:
        story.append(Paragraph("No marks recorded yet.", styles["Normal"]))
    else:
        marks_table = Table(table_data, colWidths=[4 * cm, 3 * cm, 3.5 * cm, 3 * cm])
        marks_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("TOPPADDING", (0, 0), (-1, 0), 8),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ]))
        story.append(marks_table)

    doc = SimpleDocTemplate(filename, pagesize=A4)
    doc.build(story)
    return filename

def generate_class_summary_excel(students):
    """students is a list of rows from get_all_students()."""
    os.makedirs(REPORTS_DIR, exist_ok=True)
    filename = os.path.join(REPORTS_DIR, "class_summary.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class Summary"

    headers = ["Name", "Roll No", "Class", "Section", "Attendance %", "Marks %", "Status"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_num, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for student in students:
        attendance_pct = get_attendance_percentage(student["id"])
        marks_pct = get_overall_percentage(student["id"])
        risk = predict_risk(attendance_pct, marks_pct)

        ws.append([
            student["name"], student["roll_no"], student["class_name"], student["section"],
            attendance_pct, marks_pct, risk
        ])

    for col_num, header in enumerate(headers, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = max(len(header) + 4, 14)

    wb.save(filename)
    return filename